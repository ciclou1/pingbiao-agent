"""评标报告汇编与导出（Word 底稿 / Excel 明细）。

合规要求：报告是**底稿**，AI 结论仅供参考，最终结论以评标委员会为准；
否决项必须带理由，待复核项单独列出，报告末尾留签字栏。
"""
import io
import re
from datetime import datetime, timezone

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from .models import (
    Bidder,
    Document,
    Project,
    ProjectMember,
    RequirementItem,
    ReviewResult,
    Score,
    User,
)

STATE_LABELS = {"passed": "通过", "failed": "不通过", "review": "待复核"}
TYPE_LABELS = {"qualification": "资格条件", "compliance": "符合性审查(否决项)", "scoring": "评分项"}
PROJECT_TYPES = {"goods": "货物", "engineering": "工程", "service": "服务"}
METHOD_TYPES = {"comprehensive": "综合评分法", "lowest_price": "最低价法"}


async def collect_report_data(db: AsyncSession, project_id: int) -> dict:
    project = await db.get(Project, project_id)
    if project is None:
        raise ValueError("项目不存在")

    bidders = list((await db.execute(select(Bidder).where(Bidder.project_id == project_id).order_by(Bidder.id))).scalars())
    items = list(
        (
            await db.execute(
                select(RequirementItem)
                .where(RequirementItem.project_id == project_id, RequirementItem.expert_confirmed.is_(True))
                .order_by(RequirementItem.item_type, RequirementItem.sort_order, RequirementItem.id)
            )
        ).scalars()
    )
    results = list((await db.execute(select(ReviewResult).where(ReviewResult.project_id == project_id))).scalars())
    result_map = {(r.item_id, r.bidder_id): r for r in results}

    members_rows = await db.execute(
        select(ProjectMember, User)
        .join(User, User.id == ProjectMember.user_id)
        .where(ProjectMember.project_id == project_id)
        .order_by(ProjectMember.id)
    )
    members = [
        {"name": u.display_name or u.username, "username": u.username, "role": m.member_role, "group_no": m.group_no}
        for m, u in members_rows
    ]

    doc_count = (await db.execute(select(func.count()).select_from(Document).where(Document.project_id == project_id))).scalar_one()
    by_type = {t: sum(1 for it in items if it.item_type == t) for t in TYPE_LABELS}

    # 审查矩阵
    review_rows = []
    failures, pending, decided = [], [], 0
    for it in items:
        if it.item_type == "scoring":
            continue
        cells = []
        for b in bidders:
            r = result_map.get((it.id, b.id))
            state = (r.expert_state or r.ai_state) if r else None
            if r and r.expert_state:
                decided += 1
            cells.append({"bidder_id": b.id, "state": state, "expert_state": r.expert_state if r else None})
            if state == "failed":
                failures.append(
                    {
                        "bidder": b.name,
                        "clause": it.clause_text,
                        "reason": (r.expert_reason if r and r.expert_state == "failed" else r.ai_reason) if r else "",
                        "by": "专家" if r and r.expert_state == "failed" else "AI",
                        "source_page": it.source_page,
                    }
                )
            elif state == "review":
                pending.append(
                    {"bidder": b.name, "clause": it.clause_text, "reason": r.ai_reason if r else "未审查", "source_page": it.source_page}
                )
        review_rows.append(
            {
                "clause_text": it.clause_text,
                "item_type": it.item_type,
                "category": it.category,
                "source_page": it.source_page,
                "cells": cells,
            }
        )

    # 评分汇总
    scoring_items = [it for it in items if it.item_type == "scoring"]
    score_rows = list((await db.execute(select(Score).where(Score.project_id == project_id))).scalars())
    grouped: dict[tuple[int, int], list] = {}
    for s in score_rows:
        if s.score is not None:
            grouped.setdefault((s.item_id, s.bidder_id), []).append(float(s.score))

    scores = []
    for b in bidders:
        per_item, total = [], 0.0
        for it in scoring_items:
            vals = grouped.get((it.id, b.id), [])
            avg = round(sum(vals) / len(vals), 2) if vals else None
            if avg is not None:
                total += avg
            per_item.append(
                {
                    "clause_text": it.clause_text,
                    "category": it.category,
                    "max_score": float(it.score) if it.score is not None else None,
                    "average": avg,
                    "scored_experts": len(vals),
                }
            )
        scores.append({"bidder": b.name, "items": per_item, "total_average": round(total, 2)})

    experts = [m for m in members if m["role"] in ("expert", "leader")]

    return {
        "project": {
            "id": project.id,
            "name": project.name,
            "type": PROJECT_TYPES.get(project.type, project.type),
            "method_type": METHOD_TYPES.get(project.method_type, project.method_type),
            "control_price": float(project.control_price) if project.control_price is not None else None,
            "status": project.status,
            "description": project.description,
        },
        "stats": {
            "items_total": len(items),
            "by_type": by_type,
            "bidders": len(bidders),
            "experts": len(experts),
            "documents": doc_count,
            "decided_by_expert": decided,
        },
        "bidders": [{"id": b.id, "name": b.name} for b in bidders],
        "review_rows": review_rows,
        "failures": failures,
        "failure_groups": group_failures(failures),
        "pending": pending,
        "scores": scores,
        "scoring_items_count": len(scoring_items),
        "members": members,
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


# ---------------------------------------------------------------------------
# Word 底稿
# ---------------------------------------------------------------------------

def _norm_text(t: str) -> str:
    return re.sub(r"[\s\d．.（）()、，,；;：:　\"']", "", t or "")


def group_failures(failures: list[dict], min_common: int = 6) -> list[dict]:
    """按「否决事实」归并同一投标人的多条不通过。

    真实实测：招标文件同一实质要求常在不同章节重复出现（「符合政府采购法第二十二条」
    出现在 5 个章节），一个事实（失信记录）会连锁产生 20+ 条不通过。
    逐条罗列会严重误导阅读，故按理由文本的共同片段聚类（最长公共子串 ≥ min_common 字）。
    """
    from difflib import SequenceMatcher

    groups: list[dict] = []
    for f in failures:
        reason = _norm_text(f.get("reason") or "")
        placed = False
        for g in groups:
            if g["bidder"] != f["bidder"]:
                continue
            a, b = reason, g["reason_key"]
            m = SequenceMatcher(None, a, b).find_longest_match(0, len(a), 0, len(b))
            if m.size >= min_common:
                g["items"].append(f)
                placed = True
                break
        if not placed:
            groups.append(
                {"bidder": f["bidder"], "reason_key": reason, "reason": f.get("reason") or "", "items": [f]}
            )
    for g in groups:
        g["count"] = len(g["items"])
    groups.sort(key=lambda g: (-g["count"], g["bidder"]))
    return groups


def build_docx(data: dict) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Docx()
    # 中文字体（否则 Word 里中文会走默认西文字体）
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    p = data["project"]
    title = doc.add_heading(f"{p['name']}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("评标报告（系统辅助生成底稿）")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("一、项目基本情况", level=1)
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for k, v in [
        ("项目类型", p["type"]),
        ("评标办法", p["method_type"]),
        ("控制价（万元）", p["control_price"] if p["control_price"] is not None else "—"),
        ("投标人数量", data["stats"]["bidders"]),
        ("评标委员会人数", data["stats"]["experts"]),
        ("上传文件数", data["stats"]["documents"]),
        ("项目说明", p["description"] or "—"),
    ]:
        row = info.add_row().cells
        row[0].text = str(k)
        row[1].text = str(v)
    doc.add_paragraph("")

    doc.add_heading("二、评审依据", level=1)
    s = data["stats"]
    doc.add_paragraph(
        f"经评标委员会确认的评审条目共 {s['items_total']} 条，其中资格条件 {s['by_type'].get('qualification', 0)} 条、"
        f"符合性审查（否决项）{s['by_type'].get('compliance', 0)} 条、评分项 {s['by_type'].get('scoring', 0)} 条。"
        f"以上条目由系统从招标文件自动抽取并经人工确认，作为本次评审依据。"
    )
    doc.add_paragraph(f"投标人名单：{'、'.join(b['name'] for b in data['bidders']) or '—'}")
    doc.add_paragraph(f"评标委员会成员：{'、'.join(m['name'] for m in data['members'] if m['role'] in ('expert', 'leader')) or '—'}")

    doc.add_heading("三、资格与符合性审查结果", level=1)
    if data["review_rows"]:
        headers = ["评审条目"] + [b["name"] for b in data["bidders"]]
        tbl = doc.add_table(rows=1, cols=len(headers))
        tbl.style = "Table Grid"
        for i, h in enumerate(headers):
            tbl.rows[0].cells[i].text = h
        for row in data["review_rows"]:
            cells = tbl.add_row().cells
            cells[0].text = row["clause_text"]
            by_id = {c["bidder_id"]: c for c in row["cells"]}
            for i, b in enumerate(data["bidders"]):
                cell = by_id.get(b["id"], {})
                state = cell.get("state")
                label = STATE_LABELS.get(state, "未审")
                if cell.get("expert_state"):
                    label += "（专家终值）"
                cells[i + 1].text = label
    else:
        doc.add_paragraph("（无已确认的资格/符合性条目）")

    doc.add_heading("四、否决（不通过）情况", level=1)
    if data["failures"]:
        groups = data.get("failure_groups") or []
        if any(g["count"] > 1 for g in groups):
            doc.add_paragraph(
                "说明：招标文件同一实质要求常在不同章节重复规定，同一事实可能同时违反多条条款。"
                "以下按「否决事实」归并，括号内为涉及条款条数。"
            )
        for g in groups:
            suffix = f"（涉及 {g['count']} 条条款）" if g["count"] > 1 else ""
            doc.add_paragraph(f"· {g['bidder']}：{g['reason']}{suffix}", style="List Bullet")
            if g["count"] > 1:
                for it in g["items"]:
                    doc.add_paragraph(f"{it['clause']}", style="List Bullet 2")
    else:
        doc.add_paragraph("本次审查未产生不通过结论。")

    doc.add_heading("五、待复核事项", level=1)
    if data["pending"]:
        doc.add_paragraph("以下条目材料证据不足或需核验原件，尚需评标委员会复核确认：")
        for x in data["pending"]:
            doc.add_paragraph(f"· {x['bidder']}：{x['clause']}（{x['reason']}）", style="List Bullet")
    else:
        doc.add_paragraph("无待复核事项。")

    doc.add_heading("六、评分汇总", level=1)
    if data["scoring_items_count"]:
        tbl = doc.add_table(rows=1, cols=1 + len(data["bidders"]))
        tbl.style = "Table Grid"
        tbl.rows[0].cells[0].text = "评分项"
        for i, b in enumerate(data["bidders"]):
            tbl.rows[0].cells[i + 1].text = b["name"]
        for idx, it in enumerate(data["scores"][0]["items"] if data["scores"] else []):
            cells = tbl.add_row().cells
            max_s = it["max_score"]
            cells[0].text = f"{it['clause_text']}（满分 {max_s}）" if max_s is not None else it["clause_text"]
            for i, b in enumerate(data["bidders"]):
                v = data["scores"][i]["items"][idx]["average"] if idx < len(data["scores"][i]["items"]) else None
                cells[i + 1].text = f"{v}" if v is not None else "—"
        total = tbl.add_row().cells
        total[0].text = "汇总平均总分"
        for i in range(len(data["bidders"])):
            total[i + 1].text = str(data["scores"][i]["total_average"])
    else:
        doc.add_paragraph("本项目无评分项。")

    doc.add_heading("七、评标委员会签字", level=1)
    doc.add_paragraph("本报告为系统辅助生成的评标底稿，评审结论以评标委员会审议结果为准。")
    for m in data["members"]:
        if m["role"] in ("expert", "leader"):
            doc.add_paragraph(f"{m['name']}（{m['role']}）：____________________")
    doc.add_paragraph("")
    doc.add_paragraph(f"生成时间：{data['generated_at']}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel 明细
# ---------------------------------------------------------------------------

def build_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()

    ws = wb.active
    ws.title = "审查结论明细"
    headers = ["评审条目", "类型", "分类", "招标文件页码"] + [b["name"] for b in data["bidders"]]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in data["review_rows"]:
        by_id = {c["bidder_id"]: c for c in row["cells"]}
        ws.append(
            [
                row["clause_text"],
                TYPE_LABELS.get(row["item_type"], row["item_type"]),
                row["category"],
                row["source_page"],
            ]
            + [STATE_LABELS.get(by_id.get(b["id"], {}).get("state"), "未审") for b in data["bidders"]]
        )
    ws.column_dimensions["A"].width = 60
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 16

    ws2 = wb.create_sheet("不通过与待复核")
    ws2.append(["类别", "投标人", "评审条目", "理由/说明", "来源页码"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for f in data["failures"]:
        ws2.append(["不通过", f["bidder"], f["clause"], f["reason"], f["source_page"]])
    for x in data["pending"]:
        ws2.append(["待复核", x["bidder"], x["clause"], x["reason"], x["source_page"]])
    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 60
    ws2.column_dimensions["D"].width = 50

    if data["scoring_items_count"]:
        ws3 = wb.create_sheet("评分汇总")
        ws3.append(["评分项", "满分"] + [b["name"] + " 均分" for b in data["bidders"]] + [b["name"] + " 已打分人数" for b in data["bidders"]])
        for c in ws3[1]:
            c.font = Font(bold=True)
        first = data["scores"][0]["items"] if data["scores"] else []
        for idx, it in enumerate(first):
            row = [it["clause_text"], it["max_score"]]
            row += [data["scores"][i]["items"][idx]["average"] for i in range(len(data["bidders"]))]
            row += [data["scores"][i]["items"][idx]["scored_experts"] for i in range(len(data["bidders"]))]
            ws3.append(row)
        ws3.append(["汇总平均总分", None] + [b["total_average"] for b in data["scores"]])
        ws3.column_dimensions["A"].width = 60
        for col in "BCDEFGHIJ":
            ws3.column_dimensions[col].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
