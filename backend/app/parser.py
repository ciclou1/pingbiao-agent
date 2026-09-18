"""文档解析：PyMuPDF 快速路径 + docx/xlsx 占位（M1 扩展 OCR 与专有格式）。"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field


@dataclass
class Page:
    page_no: int
    text: str
    tables: list[list[list[str]]] = field(default_factory=list)


DOCX_EXT = (".docx",)
XLSX_EXT = (".xlsx",)
PDF_EXT = (".pdf",)


def extract_pages(raw: bytes, filename: str) -> list[Page]:
    name = (filename or "").lower()
    if name.endswith(PDF_EXT):
        return _pdf_pages(raw)
    if name.endswith(DOCX_EXT):
        return _docx_pages(raw)
    if name.endswith(XLSX_EXT):
        return _xlsx_pages(raw)
    # 工程电子标书多为 zip 容器：解包可读成员（PDF/Word/XML/文本）后走正常解析
    if name.endswith(".zip") or looks_like_zip(raw):
        return _zip_pages(raw)
    return []


# 防 zip 炸弹：成员数 / 单成员大小 / 总解压大小上限
ZIP_MAX_MEMBERS = 200
ZIP_MAX_MEMBER_BYTES = 60 * 1024 * 1024
ZIP_MAX_TOTAL_BYTES = 200 * 1024 * 1024
ZIP_TEXT_EXT = (".xml", ".txt", ".csv", ".json", ".md")


def _zip_pages(raw: bytes) -> list[Page]:
    """解包 zip 容器内的可读成员。

    注意：拼出的页码是容器内顺序序号，不是原文件页码；每页开头标注来源成员名，
    证据引用据此定位到「哪个文件」。
    """
    pages: list[Page] = []
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        return []
    total = 0
    with zf:
        members = [m for m in zf.infolist() if not m.is_dir() and not m.filename.startswith("__MACOSX")]
        if len(members) > ZIP_MAX_MEMBERS:
            members = members[:ZIP_MAX_MEMBERS]
        seq = 0
        for m in members:
            lower = m.filename.lower()
            if m.file_size > ZIP_MAX_MEMBER_BYTES:
                continue
            total += m.file_size
            if total > ZIP_MAX_TOTAL_BYTES:
                break
            try:
                blob = zf.read(m)
            except (zipfile.BadZipFile, RuntimeError):
                continue  # 加密成员读不出：跳过（该容器可能整体加密）
            short = m.filename.rsplit("/", 1)[-1]
            inner: list[Page] = []
            if lower.endswith(PDF_EXT):
                inner = _pdf_pages(blob)
            elif lower.endswith(DOCX_EXT):
                inner = _docx_pages(blob)
            elif lower.endswith(XLSX_EXT):
                inner = _xlsx_pages(blob)
            elif lower.endswith(ZIP_TEXT_EXT):
                text = blob.decode("utf-8", "ignore").strip()
                # XML 去标签，留可读文本
                if lower.endswith(".xml"):
                    text = re.sub(r"<[^>]+>", " ", text)
                    text = re.sub(r"\s+", " ", text).strip()
                if text:
                    inner = [Page(page_no=0, text=text)]
            for p in inner:
                seq += 1
                p.page_no = seq
                p.text = f"【文件：{short}】\n{p.text}"
                pages.append(p)
    return pages


def guess_page_count(raw: bytes, filename: str) -> int:
    return len(extract_pages(raw, filename))


def _pdf_pages(raw: bytes) -> list[Page]:
    import fitz  # PyMuPDF

    pages: list[Page] = []
    with fitz.open(stream=raw, filetype="pdf") as pdf:
        for i, page in enumerate(pdf):
            text = page.get_text("text") or ""
            pages.append(Page(page_no=i + 1, text=text.strip()))
    return pages


def _docx_pages(raw: bytes) -> list[Page]:
    from docx import Document as Docx

    doc = Docx(io.BytesIO(raw))
    chunks: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(cells)
            if line.strip(" |"):
                chunks.append(line)
    # 段落近似分页：每 40 段一「页」
    pages = []
    step = 40
    for idx in range(0, len(chunks), step):
        pages.append(Page(page_no=idx // step + 1, text="\n".join(chunks[idx : idx + step])))
    return pages


def _xlsx_pages(raw: bytes) -> list[Page]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    pages: list[Page] = []
    page_no = 0
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if rows:
            page_no += 1
            pages.append(Page(page_no=page_no, text="", tables=[rows]))
    return pages


def looks_like_zip(raw: bytes) -> bool:
    return raw[:2] == b"PK"


def tables_to_text(tables: list) -> str:
    lines = []
    for table in tables:
        for row in table:
            line = " | ".join(str(c) for c in row).strip()
            if line.strip(" |"):
                lines.append(line)
    return "\n".join(lines)


# PDF 折行会把数字拆开（「有效期6 / 0日」其实是 60 日）。只在**跨换行**时粘回：
# 若连「数字 空格 数字」也粘，XML/表格里相邻字段（100 2000）会被并成一个假数字。
_SPLIT_DIGITS = re.compile(r"(?<=\d)[ \t 　]*\n[ \t 　]*(?=[\d.])")


def join_split_digits(text: str) -> str:
    prev = None
    while prev != text:
        prev = text
        text = _SPLIT_DIGITS.sub("", text)
    return text


# OCR 会把数字/标点识别成全角（「108．00万元」「５００」），直接进数值校验会匹配不上
_FULLWIDTH = str.maketrans(
    "０１２３４５６７８９．，：；％（）－／",
    "0123456789.,:;%()-/",
)


def normalize_ocr_text(text: str) -> str:
    """OCR 文本归一：全角数字与标点转半角、去行内多余空白。"""
    t = (text or "").translate(_FULLWIDTH)
    t = re.sub(r"[ \t]+", " ", t)
    return "\n".join(line.strip() for line in t.splitlines() if line.strip())


def chunk_pages(pages: list[Page], max_chars: int = 500, overlap: int = 80) -> list[tuple[int, str]]:
    """把页文本切成检索分块，返回 [(page_no, chunk_text)]。

    页内超长时按 max_chars 切分并保留 overlap，避免证据被截断在边界。
    """
    out: list[tuple[int, str]] = []
    for p in pages:
        text = (p.text or "").strip()
        if not text and p.tables:
            text = tables_to_text(p.tables).strip()
        if not text:
            continue
        text = join_split_digits(text)
        if len(text) <= max_chars:
            out.append((p.page_no, text))
            continue
        start = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            out.append((p.page_no, text[start:end]))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return out
